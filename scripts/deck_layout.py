"""Check a deck's layout without rendering it.

Each check returns a list of problems, empty when the deck is clean.

- Every shape sits on the slide.
- Two shapes overlap only when one holds the other whole, as a card holds its text.
- Every text frame holds its text at the size given, measured with the font's own
  metrics, and every character is one its font draws.

The third needs the font files for the deck's faces. Where they are missing it is
skipped and says so, and only a render shows whether the text fits. Set
``DECK_FONT_DIR`` to a folder holding them to look there first.
"""

import os
from collections.abc import Iterator
from dataclasses import dataclass
from functools import cache
from pathlib import Path

from fontTools.ttLib import TTFont
from PIL import ImageFont
from pptx.presentation import Presentation
from pptx.shapes.base import BaseShape
from pptx.text.text import TextFrame, _Paragraph
from pptx.util import Emu, Length

__all__ = ["FONT_DIRS", "Metrics", "fit_problems", "layout_problems", "metrics"]

FONT_DIRS = tuple(
    Path(directory)
    for directory in (
        os.environ.get("DECK_FONT_DIR", ""),
        "/mnt/c/Windows/Fonts",
        "C:/Windows/Fonts",
        "/usr/share/fonts/truetype/msttcorefonts",
        str(Path.home() / ".fonts"),
    )
    if directory
)
"""Where the font files are looked for, in order."""

_FILES = {
    ("Calibri", False): ("calibri.ttf", 0),
    ("Calibri", True): ("calibrib.ttf", 0),
    ("Courier New", False): ("cour.ttf", 0),
    ("Courier New", True): ("courbd.ttf", 0),
    ("Cambria", False): ("cambria.ttc", 0),
    ("Cambria", True): ("cambriab.ttf", 0),
}
# Single line spacing as a multiple of the font size, from each face's ascent, descent
# and line gap. A line mixing faces takes the tallest. Courier New's metrics give
# 1.133, but PowerPoint sets it 1.2 apart, measured on a render.
_LINE_HEIGHT = {"Calibri": 1.22, "Courier New": 1.2, "Cambria": 1.172}
_SLACK = 0.02  # inches, for rounding in the renderer
_MEASURE_SIZE = 1000  # points, so widths come out in thousandths of an em


@dataclass(frozen=True)
class Metrics:
    """Text widths in the deck's faces, read from their font files."""

    directory: Path

    def width(self, text: str, font: str, size: float, *, bold: bool) -> float:
        """How wide ``text`` sets, in inches, at ``size`` points."""
        name, index = _FILES[(font, bold)]
        face = _face(self.directory / name, index)
        return face.getlength(text) * size / _MEASURE_SIZE / 72

    def lacks(self, text: str, font: str, *, bold: bool) -> str:
        """The characters of ``text`` the face has no glyph for, in order."""
        name, index = _FILES[(font, bold)]
        drawn = _characters(self.directory / name, index)
        return "".join(dict.fromkeys(c for c in text if ord(c) not in drawn))


@cache
def _face(path: Path, index: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), _MEASURE_SIZE, index=index)


@cache
def _characters(path: Path, index: int) -> frozenset[int]:
    return frozenset(TTFont(str(path), fontNumber=index).getBestCmap() or {})


def metrics() -> Metrics | None:
    """The font metrics, or ``None`` if no directory holds every face's file."""
    for directory in FONT_DIRS:
        if all((directory / name).is_file() for name, _ in _FILES.values()):
            return Metrics(directory)
    return None


def _inches(length: Length | int | None) -> float:
    return Emu(length or 0).inches


Bounds = tuple[float, float, float, float]


def _bounds(shape: BaseShape) -> Bounds:
    left, top = _inches(shape.left), _inches(shape.top)
    return left, top, left + _inches(shape.width), top + _inches(shape.height)


def _holds(outer: Bounds, inner: Bounds) -> bool:
    return (
        outer[0] <= inner[0] + 1e-6
        and outer[1] <= inner[1] + 1e-6
        and outer[2] >= inner[2] - 1e-6
        and outer[3] >= inner[3] - 1e-6
    )


def _overlap(first: Bounds, second: Bounds) -> bool:
    return min(first[2], second[2]) > max(first[0], second[0]) and min(
        first[3], second[3]
    ) > max(first[1], second[1])


def _slides(presentation: Presentation) -> Iterator[tuple[int, list[BaseShape]]]:
    for number, slide in enumerate(presentation.slides, start=1):
        yield number, list(slide.shapes)


def layout_problems(presentation: Presentation) -> list[str]:
    """Shapes off the slide, and shapes that overlap without one holding the other."""
    width = _inches(presentation.slide_width)
    height = _inches(presentation.slide_height)
    problems = []
    for number, shapes in _slides(presentation):
        bounds = [_bounds(shape) for shape in shapes]
        for shape, (left, top, right, bottom) in zip(shapes, bounds, strict=True):
            if left < 0 or top < 0 or right > width or bottom > height:
                problems.append(f"slide {number}: {shape.name} runs off the slide")
        for i, first in enumerate(bounds):
            for j in range(i + 1, len(bounds)):
                second = bounds[j]
                held = _holds(first, second) or _holds(second, first)
                if _overlap(first, second) and not held:
                    problems.append(
                        f"slide {number}: {shapes[i].name} and {shapes[j].name} overlap"
                    )
    return problems


@dataclass(frozen=True)
class _Word:
    """A word and the space after it, in its run's style."""

    text: str
    space: str
    font: str
    size: float
    bold: bool


def _words(paragraph: _Paragraph) -> list[_Word]:
    words = []
    for run in paragraph.runs:
        if not run.text or run.font.size is None:
            continue
        font, bold = run.font.name or "Calibri", bool(run.font.bold)
        pieces = run.text.split(" ")
        for index, piece in enumerate(pieces):
            space = " " if index < len(pieces) - 1 else ""
            words.append(_Word(piece, space, font, run.font.size.pt, bold))
    return words


def _set_lines(
    paragraph: _Paragraph, available: float, wrap: bool, measure: Metrics
) -> tuple[list[float], float]:
    """The width of each line a paragraph sets in, and its largest size."""
    words = _words(paragraph)
    if not words:
        return [0.0], 0.0

    def width(word: _Word, *, spaced: bool) -> float:
        text = word.text + (word.space if spaced else "")
        return measure.width(text, word.font, word.size, bold=word.bold)

    # A line's width stops at its last word. The space after it hangs past the edge.
    lines, current, space = [], 0.0, 0.0
    for word in words:
        bare = width(word, spaced=False)
        if wrap and current and current + space + bare > available:
            lines.append(current)
            current, space = 0.0, 0.0
        current += space + bare
        space = width(word, spaced=True) - bare
    lines.append(current)
    return lines, max(word.size for word in words)


def _frame_problems(
    frame: TextFrame, box: tuple[float, float], measure: Metrics
) -> Iterator[str]:
    """Why ``frame`` does not hold its text, or nothing."""
    available_width = box[0] - _inches(frame.margin_left) - _inches(frame.margin_right)
    available_height = box[1] - _inches(frame.margin_top) - _inches(frame.margin_bottom)
    wrap = frame.word_wrap is not False
    needed = 0.0
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            font = run.font.name or "Calibri"
            missing = measure.lacks(run.text, font, bold=bool(run.font.bold))
            if missing:
                yield f"{font} has no glyph for {missing!r}"
        lines, size = _set_lines(paragraph, available_width, wrap, measure)
        if max(lines) > available_width + _SLACK:
            yield f"a line is {max(lines):.2f} in wide in {available_width:.2f} in"
        faces = {run.font.name or "Calibri" for run in paragraph.runs} or {"Calibri"}
        spacing = paragraph.line_spacing
        multiple = spacing if isinstance(spacing, float) else 1.0
        line_height = max(_LINE_HEIGHT[face] for face in faces)
        needed += len(lines) * size * line_height * multiple / 72
        if paragraph.space_after is not None:
            needed += paragraph.space_after.pt / 72
    if needed > available_height + _SLACK:
        yield f"the text needs {needed:.2f} in of {available_height:.2f} in"


def fit_problems(presentation: Presentation, measure: Metrics) -> list[str]:
    """Text frames whose text runs past their box at the size given."""
    problems = []
    for number, shapes in _slides(presentation):
        for shape in shapes:
            frame = getattr(shape, "text_frame", None) if shape.has_text_frame else None
            if frame is None:
                continue
            box = (_inches(shape.width), _inches(shape.height))
            problems.extend(
                f"slide {number}: {shape.name}: {problem}"
                for problem in _frame_problems(frame, box, measure)
            )
    return problems
