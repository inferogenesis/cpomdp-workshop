"""Drawing primitives for the workshop deck, in the light palette of the route frames.

Every position and size is in inches on a 16:9 slide, 13.333 wide and 7.5 high. A
page is one slide with its title in the same place as every other page's. Code sits
in a filled panel, printed output in an outlined one. A placeholder the owner has to
fill is drawn dashed in amber, so it cannot pass for finished content.

Shapes are named by kind (``title``, ``code``, ``printed``, ``text``, ``picture``,
``placeholder``, ``card``, ``pill``, ``mark``, ``track``, ``bar``). The layout check
and the deck test read the names, so a number inside program output is told apart
from a number in prose.
"""

import math
import re
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.oxml import parse_xml
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pptx.util import Emu, Inches, Pt

from cpomdp_workshop import plot

__all__ = [
    "AMBER",
    "AMBER_MARK",
    "BACKGROUND",
    "BLUE",
    "CODE_FONT",
    "CONTENT_TOP",
    "GREEN",
    "INK",
    "MARGIN",
    "MATH_FONT",
    "MUTED",
    "PANEL",
    "RULE",
    "SLIDE_HEIGHT",
    "SLIDE_WIDTH",
    "TEXT_FONT",
    "Deck",
    "Page",
    "Paragraph",
    "Run",
    "code_height",
    "code_size",
    "without_imports",
]

# The route frames' palette, so a slide and the figure on it agree. The accents are
# the light theme of Inferogenesis's site: a darker amber, blue and green for text,
# which the lighter figure colours would leave too faint, and the bright amber of its
# mark for outlines and the timer's bar.
BACKGROUND = plot.LIGHT.background.lstrip("#")
PANEL = plot.LIGHT.panel.lstrip("#")
RULE = plot.LIGHT.rule.lstrip("#")
INK = plot.LIGHT.ink.lstrip("#")
MUTED = plot.LIGHT.muted.lstrip("#")
AMBER = "8A6000"
AMBER_MARK = "E69F00"
BLUE = "005F94"
GREEN = "007A5A"

TEXT_FONT = "Calibri"
CODE_FONT = "Courier New"
MATH_FONT = "Cambria"
# Courier New has no superscript plus or minus, and cpomdp's messages print both (μ⁺,
# μ⁻). They are set in the math face, which has them, so no renderer has to guess a
# substitute. Cambria Math has them too, with a line height three times the code face's.
_CODE_FALLBACK = re.compile("([⁺⁻])")

SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5
MARGIN = 0.6
CONTENT_TOP = 1.45

_TITLE = (MARGIN, 0.42, 10.2, 0.8)  # x, y, w, h
_SEGMENT = (11.03, 0.58, 1.7, 0.42)
_PANEL_INSET = 0.18
_CODE_ADVANCE = 0.6  # Courier New sets every glyph 0.6 em wide
_CODE_LINE = 1.2  # and PowerPoint sets its lines 1.2 em apart, measured on a render
_SLOT = re.compile(r"^\s*\w+ = \.\.\.")


@dataclass(frozen=True)
class Run:
    """A stretch of text in one style inside a paragraph.

    Attributes:
        text: the characters.
        colour: hex without the hash, or ``None`` for the paragraph's colour.
        bold: whether it is bold.
        font: the typeface, or ``None`` for the paragraph's.
        size: the size in points, or ``None`` for the paragraph's.
        link: a URL the text opens when clicked, or ``None``.
    """

    text: str
    colour: str | None = None
    bold: bool = False
    font: str | None = None
    size: float | None = None
    link: str | None = None


Paragraph = str | Sequence[Run]


def _rgb(colour: str) -> RGBColor:
    return RGBColor.from_string(colour)


def code_size(
    source: str, width: float, height: float, *, largest: float = 14.0
) -> float:
    """The largest size, in half points, at which ``source`` fits a code panel.

    Args:
        source: the code.
        width: the panel's width, in inches.
        height: the most height the panel may take, in inches.
        largest: the size not to go past.

    Returns:
        The size in points.
    """
    lines = source.splitlines()
    inner_width = width - 2 * _PANEL_INSET
    inner_height = height - 1.4 * _PANEL_INSET
    by_width = inner_width / (max(map(len, lines)) * _CODE_ADVANCE / 72)
    by_height = inner_height / (len(lines) * _CODE_LINE / 72)
    return min(largest, math.floor(2 * min(by_width, by_height)) / 2)


def code_height(source: str, size: float) -> float:
    """The height, in inches, of a code panel that holds ``source`` at ``size``."""
    return len(source.splitlines()) * size * _CODE_LINE / 72 + 1.4 * _PANEL_INSET + 0.04


def without_imports(source: str) -> str:
    """``source`` with its leading import block and the blank lines after it cut.

    A slide shows the lines a learner reads. The imports are in the cell.
    """
    lines = source.splitlines()
    start, in_parentheses = 0, False
    for index, line in enumerate(lines):
        stripped = line.strip()
        if in_parentheses:
            in_parentheses = not stripped.startswith(")")
        elif stripped.startswith(("import ", "from ")):
            in_parentheses = stripped.endswith("(")
        elif stripped:
            start = index
            break
    return "\n".join(lines[start:])


class Page:
    """One slide, and the calls that put things on it."""

    def __init__(self, slide: Slide) -> None:
        self.slide = slide

    def _box(self, x: float, y: float, w: float, h: float, name: str) -> TextFrame:
        shape = self.slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.name = name
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.0)
        frame.margin_top = frame.margin_bottom = Inches(0.0)
        return frame

    def _panel(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        name: str,
        *,
        fill: str | None,
        line: str | None,
        dashed: bool = False,
    ) -> Shape:
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.name = name
        shape.adjustments[0] = min(0.08, 0.12 / min(w, h))
        shape.shadow.inherit = False
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = _rgb(line)
            shape.line.width = Pt(1.25 if dashed else 1)
            if dashed:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        frame = shape.text_frame
        frame.word_wrap = True
        inset = Inches(_PANEL_INSET)
        frame.margin_left = frame.margin_right = inset
        frame.margin_top = frame.margin_bottom = Inches(_PANEL_INSET * 0.7)
        return shape

    @staticmethod
    def _fill(
        frame: TextFrame,
        paragraphs: Sequence[Paragraph],
        *,
        size: float,
        colour: str,
        font: str,
        bold: bool = False,
        align: PP_ALIGN = PP_ALIGN.LEFT,
        space_after: float = 0.0,
        line_spacing: float = 1.0,
    ) -> None:
        for index, paragraph in enumerate(paragraphs):
            target = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            target.alignment = align
            target.line_spacing = line_spacing
            if index < len(paragraphs) - 1:
                target.space_after = Pt(space_after)
            runs = [Run(paragraph)] if isinstance(paragraph, str) else paragraph
            for piece in runs:
                face = piece.font or font
                for text, fallback in _glyph_runs(piece.text, face):
                    run = target.add_run()
                    run.text = text
                    run.font.size = Pt(piece.size or size)
                    run.font.name = MATH_FONT if fallback else face
                    run.font.bold = piece.bold or bold
                    run.font.color.rgb = _rgb(piece.colour or colour)
                    if piece.link is not None:
                        run.hyperlink.address = piece.link

    def title(self, text: str, segment: str) -> None:
        """The slide's title, and the part of the session it belongs to, top right."""
        frame = self._box(*_TITLE, name="title")
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        self._fill(frame, [text], size=32, colour=INK, font=TEXT_FONT, bold=True)
        self.pill(*_SEGMENT, segment, colour=MUTED, size=13)

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        paragraphs: Sequence[Paragraph],
        *,
        size: float = 18,
        colour: str = INK,
        bold: bool = False,
        font: str = TEXT_FONT,
        align: PP_ALIGN = PP_ALIGN.LEFT,
        anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
        space_after: float = 8,
        name: str = "text",
    ) -> None:
        """Prose, one paragraph per entry."""
        frame = self._box(x, y, w, h, name)
        frame.vertical_anchor = anchor
        self._fill(
            frame, paragraphs, size=size, colour=colour, font=font, bold=bold,
            align=align, space_after=space_after,
        )  # fmt: skip

    def code(
        self, x: float, y: float, w: float, h: float, source: str, *, size: float = 12
    ) -> None:
        """Python in a filled panel. Lines to write are amber, comments muted."""
        shape = self._panel(x, y, w, h, "code", fill=PANEL, line=RULE)
        shape.text_frame.word_wrap = False
        shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        self._fill(
            shape.text_frame, [_code_line(line) for line in source.splitlines()],
            size=size, colour=INK, font=CODE_FONT,
        )  # fmt: skip

    def printed(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        output: str,
        *,
        size: float = 12,
        wrap: bool = False,
    ) -> None:
        """What a cell printed, verbatim, in an outlined panel."""
        shape = self._panel(x, y, w, h, "printed", fill=None, line=RULE)
        shape.text_frame.word_wrap = wrap
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        self._fill(
            shape.text_frame, output.splitlines(), size=size, colour=INK, font=CODE_FONT
        )

    def card(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        paragraphs: Sequence[Paragraph],
        *,
        size: float = 18,
        colour: str = INK,
        font: str = TEXT_FONT,
        anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
        space_after: float = 8,
        line: str | None = None,
    ) -> None:
        """Text on a filled panel."""
        shape = self._panel(x, y, w, h, "card", fill=PANEL, line=line)
        shape.text_frame.vertical_anchor = anchor
        self._fill(
            shape.text_frame, paragraphs, size=size, colour=colour, font=font,
            space_after=space_after,
        )  # fmt: skip

    def placeholder(
        self, x: float, y: float, w: float, h: float, label: str, *, size: float = 16
    ) -> None:
        """A dashed amber box naming what the owner still has to put here."""
        shape = self._panel(
            x, y, w, h, "placeholder", fill=None, line=AMBER_MARK, dashed=True
        )
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        self._fill(
            shape.text_frame, [label], size=size, colour=AMBER, font=TEXT_FONT,
            align=PP_ALIGN.CENTER,
        )  # fmt: skip

    def pill(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        label: str,
        *,
        colour: str = AMBER,
        size: float = 14,
        font: str = TEXT_FONT,
        line: str = RULE,
    ) -> None:
        """A short label in a rounded outline."""
        shape = self._panel(x, y, w, h, "pill", fill=None, line=line)
        frame = shape.text_frame
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = Inches(0.08)
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        self._fill(
            frame, [label], size=size, colour=colour, font=font, align=PP_ALIGN.CENTER
        )

    def picture(self, path: Path, x: float, y: float, w: float) -> float:
        """A figure from ``assets/`` at width ``w``. Returns its height."""
        shape = self.slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), Inches(w)
        )
        shape.name = f"picture {path.name}"
        return Emu(shape.height).inches

    def link(self, x: float, y: float, w: float, h: float, url: str, *, size: float):
        """A URL in the code face that opens when clicked."""
        frame = self._box(x, y, w, h, "link")
        frame.word_wrap = False
        run = frame.paragraphs[0].add_run()
        run.text = url.removeprefix("https://")
        run.font.size = Pt(size)
        run.font.name = CODE_FONT
        run.font.color.rgb = _rgb(INK)
        run.hyperlink.address = url

    def block(
        self, x: float, y: float, w: float, h: float, kind: str, *, label: str = ""
    ) -> None:
        """One stretch of the agenda's bar.

        ``"talk"`` is a plain panel, ``"code"`` is filled in the mark's amber and
        ``"open"`` is dashed, like time not yet given to anything.
        """
        fill, line = {
            "talk": (PANEL, RULE),
            "code": (AMBER_MARK, None),
            "open": (None, MUTED),
        }[kind]
        shape = self._panel(
            x, y, w, h, "block", fill=fill, line=line, dashed=kind == "open"
        )
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        shape.text_frame.margin_left = shape.text_frame.margin_right = Inches(0.02)
        if label:
            self._fill(
                shape.text_frame, [label], size=14, colour=INK, font=TEXT_FONT,
                align=PP_ALIGN.CENTER,
            )  # fmt: skip

    def mark(self, x: float, y: float, height: float) -> None:
        """Inferogenesis's mark: an amber ring round a dot, over a stem.

        Drawn from the site's 100 by 100 SVG, so it scales without an image file.
        """
        unit = height / 80  # the mark spans y = 8 to 88 of its 100
        group = self.slide.shapes.add_group_shape()
        group.name = "mark"
        ring = 6.5 * unit
        parts = (
            (MSO_SHAPE.OVAL, 35, 11, 30, 30, None, AMBER_MARK),
            (MSO_SHAPE.OVAL, 45.5, 21.5, 9, 9, INK, None),
            (MSO_SHAPE.ROUNDED_RECTANGLE, 43.5, 44, 13, 44, INK, None),
        )
        for kind, left, top, width, tall, fill, line in parts:
            shape = group.shapes.add_shape(
                kind,
                Inches(x + (left - 32) * unit), Inches(y + (top - 8) * unit),
                Inches(width * unit), Inches(tall * unit),
            )  # fmt: skip
            shape.shadow.inherit = False
            if fill is None:
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(fill)
            if line is None:
                shape.line.fill.background()
            else:
                shape.line.color.rgb = _rgb(line)
                shape.line.width = Inches(ring)
            if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                shape.adjustments[0] = 0.5

    def countdown(
        self,
        x: float,
        y: float,
        w: float,
        minutes: int,
        *,
        then: Sequence[Paragraph],
        size: float = 22,
    ) -> None:
        """A bar that fills over ``minutes`` from when the slide shows, then ``then``.

        The bar fills left to right with no click, marked every five minutes. When it
        is full, ``then`` appears under the marks. Leaving the slide starts the bar
        again on return, so the slide stays up while the room works.
        """
        track = self._panel(x, y, w, 0.7, "track", fill=PANEL, line=RULE)
        track.adjustments[0] = 0.5
        bar = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x + 0.06), Inches(y + 0.06), Inches(w - 0.12), Inches(0.58),
        )  # fmt: skip
        bar.name = "bar"
        bar.adjustments[0] = 0.5
        bar.shadow.inherit = False
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(AMBER_MARK)
        bar.line.fill.background()
        marks = range(0, minutes + 1, 5)
        for minute in marks:
            label = f"{minute} min" if minute == minutes else str(minute)
            centre = x + w * minute / minutes
            self.text(
                min(max(centre - 0.6, x), x + w - 1.2), y + 0.85, 1.2, 0.45, [label],
                size=16, colour=MUTED, space_after=0,
                align=PP_ALIGN.LEFT if minute == 0 else
                PP_ALIGN.RIGHT if minute == minutes else PP_ALIGN.CENTER,
            )  # fmt: skip
        self.text(x, y + 1.5, w, 0.6, then, size=size, bold=True, space_after=0)
        reveal = self.slide.shapes[-1]
        self.slide._element.append(
            _timing(
                (
                    _Effect(bar.shape_id, "wipe", minutes * 60_000),
                    _Effect(reveal.shape_id, "appear", 1),
                )
            )
        )

    def notes(self, text: str) -> None:
        """The speaker notes."""
        frame = self.slide.notes_slide.notes_text_frame
        if frame is None:  # the default notes master always has a body
            raise ValueError("this slide's notes page has no body to write in")
        frame.text = text


@dataclass(frozen=True)
class _Effect:
    """One entrance in a slide's timeline.

    Attributes:
        shape_id: the shape it brings in.
        kind: ``"wipe"`` from the left, or ``"appear"`` at once.
        duration: how long it takes, in milliseconds.
    """

    shape_id: int
    kind: str
    duration: int


_PRESENTATION = "http://schemas.openxmlformats.org/presentationml/2006/main"
_PRESETS = {"wipe": ("22", "8"), "appear": ("1", "0")}  # PowerPoint's ID, subtype


def _timing(effects: Sequence[_Effect]):
    """A slide's ``p:timing``: each effect after the last, the first with the slide.

    PowerPoint writes "after previous" this way: one group the slide's start
    triggers, holding one delayed step per effect.
    """
    ids = iter(range(4, 1000))
    steps, start = [], 0
    for effect in effects:
        preset, subtype = _PRESETS[effect.kind]
        target = f'<p:tgtEl><p:spTgt spid="{effect.shape_id}"/></p:tgtEl>'
        wipe = (
            f'<p:animEffect transition="in" filter="wipe(left)"><p:cBhvr>'
            f'<p:cTn id="{next(ids)}" dur="{effect.duration}"/>{target}'
            "</p:cBhvr></p:animEffect>"
            if effect.kind == "wipe"
            else ""
        )
        step, node, shown = next(ids), next(ids), next(ids)
        steps.append(
            f'<p:par><p:cTn id="{step}" fill="hold"><p:stCondLst>'
            f'<p:cond delay="{start}"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{node}" presetID="{preset}" presetClass="entr" '
            f'presetSubtype="{subtype}" fill="hold" grpId="0" nodeType="afterEffect">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{shown}" dur="1" fill="hold"><p:stCondLst>'
            f'<p:cond delay="0"/></p:stCondLst></p:cTn>{target}<p:attrNameLst>'
            "<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
            '<p:to><p:strVal val="visible"/></p:to></p:set>'
            f"{wipe}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
        )
        start += effect.duration
    builds = "".join(
        f'<p:bldP spid="{effect.shape_id}" grpId="0" animBg="1"/>' for effect in effects
    )
    return parse_xml(
        f'<p:timing xmlns:p="{_PRESENTATION}"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
        '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst>'
        f"<p:childTnLst>{''.join(steps)}</p:childTnLst></p:cTn></p:par>"
        "</p:childTnLst></p:cTn>"
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/>'
        "</p:tgtEl></p:cond></p:prevCondLst>"
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/>'
        "</p:tgtEl></p:cond></p:nextCondLst>"
        "</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>"
        f"<p:bldLst>{builds}</p:bldLst></p:timing>"
    )


def _glyph_runs(text: str, face: str) -> list[tuple[str, bool]]:
    """``text`` split where ``face`` lacks a glyph, each piece flagged if it does."""
    if face != CODE_FONT:
        return [(text, False)]
    pieces = _CODE_FALLBACK.split(text)
    return [(piece, index % 2 == 1) for index, piece in enumerate(pieces) if piece]


def _code_line(line: str) -> list[Run]:
    """One line of code as runs: amber if the learner writes it, comments muted."""
    if _SLOT.match(line):
        return [Run(line, colour=AMBER, bold=True)]
    if line.lstrip().startswith("#"):
        return [Run(line, colour=MUTED)]
    code, hash_mark, comment = line.partition("  # ")
    if not hash_mark:
        return [Run(line or " ")]
    return [Run(code), Run(f"{hash_mark}{comment}", colour=MUTED)]


def _theme_links_in(presentation, colour: str) -> None:
    """Set the theme's link colours, visited or not, to ``colour``.

    PowerPoint draws a hyperlink in the theme's link colour whatever its run says.
    The deck draws its links in ink, as it did on the dark background where the blank
    template's dark blue vanished.
    """
    theme = presentation.slide_master.part.part_related_by(RELATIONSHIP_TYPE.THEME)
    blob = theme.blob
    for name in ("hlink", "folHlink"):
        slot = rf"<a:{name}>.*?</a:{name}>".encode()
        recoloured = f'<a:{name}><a:srgbClr val="{colour}"/></a:{name}>'.encode()
        blob, found = re.subn(slot, recoloured, blob, flags=re.DOTALL)
        if found != 1:
            raise ValueError(f"the template's theme has {found} {name} colours, not 1")
    theme._blob = blob


class Deck:
    """A 16:9 deck on the light background of Inferogenesis's site."""

    def __init__(self, title: str, author: str) -> None:
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(SLIDE_WIDTH)
        self.presentation.slide_height = Inches(SLIDE_HEIGHT)
        # The blank template names its own author. The deck names its own.
        properties = self.presentation.core_properties
        properties.title = title
        properties.author = properties.last_modified_by = author
        self._blank = self.presentation.slide_layouts[6]
        _theme_links_in(self.presentation, INK)

    def page(self, title: str | None = None, segment: str = "") -> Page:
        """A new slide on the background, titled unless ``title`` is ``None``."""
        slide = self.presentation.slides.add_slide(self._blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(BACKGROUND)
        page = Page(slide)
        if title is not None:
            page.title(title, segment)
        return page

    def save(self, path: Path) -> None:
        """Write the deck to ``path``."""
        path.parent.mkdir(parents=True, exist_ok=True)
        self.presentation.save(str(path))
