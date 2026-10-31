import hashlib
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import jax.numpy as jnp
import pytest

pytest.importorskip("pptx", reason="the deck extra builds the slides")

from cpomdp import (
    Agent,
    Belief,
    CouplingGraph,
    CouplingGraphBackend,
    FixedSensor,
    GaussianTransition,
    KalmanBackend,
)
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE

import build_deck
import deck_layout
import lesson
from cpomdp_workshop import LAUNCH, model_with, pad_goal, plot
from cpomdp_workshop.check import backend_differences
from cpomdp_workshop.solutions import CLEAR_SKY_ROCK
from deck_kit import without_imports

PROSE = ("title", "text", "card", "pill", "placeholder")
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"
PRESENTATION = "http://schemas.openxmlformats.org/presentationml/2006/main"
NUMBER = re.compile(r"(?<![\w.])\d+(?:\.\d+)*(?:e-?\d+)?")
CLOCK = re.compile(r"\b\d+:\d\d\b")

MEASURED = {
    "0.000000": "test_no_move_of_the_1960_rocket_is_worth_more_than_any_other",
    "0.13": "test_the_1960_rocket_reports_where_it_lands",
    "1.41": "test_the_beacon_sensor_bends_the_path_to_the_beacon_and_lands",
    "0.24": "test_the_beacon_sensor_bends_the_path_to_the_beacon_and_lands",
    "1.49": "test_a_rocket_too_curious_to_land_stays_short_however_far_it_looks",
    "1.65": "test_a_rocket_too_curious_to_land_stays_short_however_far_it_looks",
    "1.86": "test_a_rocket_too_curious_to_land_stays_short_however_far_it_looks",
    "6": "test_the_1960_rocket_hits_the_rock_on_step_six_whatever_the_noise_does",
    "1.71": "test_the_default_slider_rocket_clears_the_rock_and_lands",
    "4.83": "test_the_default_slider_rocket_clears_the_rock_and_lands",
    "0.09": "test_the_default_slider_rocket_clears_the_rock_and_lands",
    "1.91": "test_the_k_point_two_rocket_clears_the_rock_and_lands",
    "5.18": "test_the_k_point_two_rocket_clears_the_rock_and_lands",
    "0.15": "test_the_k_point_two_rocket_clears_the_rock_and_lands",
    "5": "test_in_clear_sky_the_rock_is_hit_on_the_curious_rockets_routes",
    "9": "test_the_search_is_decided_rather_than_sampled",
    "81": "test_two_moves_ahead_scores_eighty_one_plans_a_step",
    "78125": "test_the_coda_certificate_covers_every_seven_step_axis_plan",
    "1e-9": "test_the_kalman_filter_and_the_one_node_graph_agree",
    "0.5.0": "test_the_numbers_were_measured_on_the_pinned_release",
}
"""Each measured number in the prose or the notes, and the claim test asserting it."""

SET = {
    "0.08": "k, the slider's default",
    "0.20": "k, the second curious rocket",
    "0.4": "k, task 2's reference",
    "0.40": "k, task 2's reference",
    "0.6": "k, too curious to land",
    "0.5": "the compass step",
    "4.2": "the clear-sky rock's x",
    "0.7": "the clear-sky rock's y, below the line",
    "0": "the first seed, and the start of a timer's bar",
    "1": "a task, poll, checkpoint, seed, step or answer number",
    "2": "a task, poll, checkpoint, seed, step or answer number",
    "3": "a task, poll, checkpoint, seed, step or answer number",
    "4": "an answer number",
    "10": "the seeds a bonus flies, and a timer's mark",
    "15": "the minutes each task codes for, and the minutes left open",
    "45": "the minutes of code in all",
    "60": "a poll's length in seconds",
    "48": "the hours before the recording is posted",
    "1960": "the year of the controller",
    "2021": "the year of the reduction to LQR",
    "2026": "the year of the symposium",
    "2.2": "a figure number in the Active Inference book",
    "2.6": "an equation number in the Active Inference book",
    "3.11": "the oldest Python the workshop runs on",
}
"""Numbers the deck sets or names rather than measures."""


@pytest.fixture(scope="module")
def deck():
    return build_deck.build().presentation


def slide_of(deck, slide):
    return deck.slides[build_deck.number_of(slide) - 1]


def prose_and_notes(presentation):
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.name.split()[0] in PROSE:
                yield shape.text_frame.text
        yield slide.notes_slide.notes_text_frame.text


def pictures(slide):
    return [shape for shape in slide.shapes if shape.name.startswith("picture")]


@pytest.mark.slow
def test_the_deck_has_thirty_one_slides_each_with_notes(deck):
    assert len(deck.slides) == 31
    for slide in deck.slides:
        assert slide.notes_slide.notes_text_frame.text.strip()


def minutes_of(kind):
    return sum(
        length
        for block in build_deck.AGENDA
        for part, length in block.parts
        if part == kind
    )


def test_the_two_hours_are_an_hour_of_talk_45_of_code_and_15_open():
    assert (minutes_of("talk"), minutes_of("code"), minutes_of("open")) == (60, 45, 15)


def test_each_task_codes_for_as_long_as_its_timer_runs():
    tasks = [block for block in build_deck.AGENDA if block.label.startswith("Task")]
    assert [block.label for block in tasks] == ["Task 1", "Task 2", "Task 3"]
    for block in tasks:
        assert [kind for kind, _ in block.parts] == ["talk", "code", "talk"]
        assert block.parts[1][1] == build_deck.TASK_MINUTES == 15


def test_the_clock_reads_each_start_off_the_agenda():
    assert build_deck.clock("Welcome") == "0:00"
    assert build_deck.clock("Task 1", 1) == "0:16"
    assert build_deck.clock("Questions") == "1:45"
    with pytest.raises(ValueError, match="no block called 'Break'"):
        build_deck.clock("Break")


@pytest.mark.slow
@pytest.mark.parametrize(
    "timer",
    [build_deck.task_1_timer, build_deck.task_2_timer, build_deck.task_3_timer],
)
def test_a_timer_fills_its_bar_over_fifteen_minutes_with_no_click(deck, timer):
    slide = slide_of(deck, timer)
    (bar,) = [shape for shape in slide.shapes if shape.name == "bar"]
    timing = slide._element.find(f"{{{PRESENTATION}}}timing")
    assert timing is not None
    wipe = timing.find(f".//{{{PRESENTATION}}}animEffect")
    assert wipe is not None
    assert wipe.get("filter") == "wipe(left)"
    behaviour = wipe.find(f"{{{PRESENTATION}}}cBhvr")
    assert behaviour is not None
    duration = behaviour.find(f"{{{PRESENTATION}}}cTn")
    target = behaviour.find(f".//{{{PRESENTATION}}}spTgt")
    assert duration is not None
    assert target is not None
    assert duration.get("dur") == str(15 * 60_000)
    assert target.get("spid") == str(bar.shape_id)
    starts = timing.findall(f".//{{{PRESENTATION}}}cond[@evt='onBegin']")
    assert len(starts) == 1, "the effects start with the slide, not on a click"


def test_every_slide_is_on_the_light_palette(deck):
    for slide in deck.slides:
        assert str(slide.background.fill.fore_color.rgb) == plot.LIGHT.background[1:]


@pytest.mark.slow
def test_no_shape_runs_off_a_slide_or_overlaps_another(deck):
    assert deck_layout.layout_problems(deck) == []


@pytest.mark.slow
def test_every_text_frame_holds_its_text(deck):
    measure = deck_layout.metrics()
    if measure is None:
        pytest.skip("no font files to measure with")
    assert deck_layout.fit_problems(deck, measure) == []


@pytest.mark.slow
def test_the_route_frames_build_up_in_one_place(deck):
    route = {
        build_deck.task_1_result: "straight",
        build_deck.task_2_result: "beacon",
        build_deck.task_3_blackout: "blackout",
        build_deck.task_3_rock: "rock",
    }
    places = set()
    for slide, stage in route.items():
        (picture,) = pictures(slide_of(deck, slide))
        assert picture.name.endswith(f"_{stage}.png")
        places.add((picture.left, picture.top, picture.width, picture.height))
    assert len(places) == 1


@pytest.mark.slow
@pytest.mark.parametrize(
    ("slide", "session"),
    [
        (build_deck.task_1_code, lesson.SESSION_1),
        (build_deck.task_2_code, lesson.SESSION_2),
        (build_deck.task_3_code, lesson.SESSION_3),
    ],
)
def test_a_task_slide_shows_the_lessons_own_code_with_both_lines_to_write(
    deck, slide, session
):
    codes = [s for s in slide_of(deck, slide).shapes if s.name == "code"]
    shown = codes[0].text_frame.text.replace("\v", "\n").splitlines()
    shown = [line.rstrip() for line in shown]
    cell = without_imports(session.attempt).splitlines()
    assert shown == cell[: len(shown)]
    slots = [line for line in cell if re.match(r"^\s*\w+ = \.\.\.", line)]
    assert len(slots) == 2
    assert all(slot in shown for slot in slots)


def test_the_graph_slide_code_flies_the_same_beliefs_as_the_kalman_filter():
    namespace: dict[str, Any] = {
        "jnp": jnp,
        "Agent": Agent,
        "Belief": Belief,
        "CouplingGraph": CouplingGraph,
        "CouplingGraphBackend": CouplingGraphBackend,
        "FixedSensor": FixedSensor,
        "GaussianTransition": GaussianTransition,
        "LAUNCH": LAUNCH,
        "goal": pad_goal(),
    }
    exec(build_deck.GRAPH_CODE, namespace)
    assert isinstance(namespace["rocket"], Agent)
    (difference,) = backend_differences(
        KalmanBackend(model_with()), namespace["backend"]
    )
    assert difference < 1e-9


def test_the_clear_sky_rock_is_where_the_slide_says():
    (x, y), _ = CLEAR_SKY_ROCK
    assert (x, y) == (4.2, -0.7)


@pytest.mark.slow
def test_every_number_in_the_prose_and_notes_is_measured_or_set(deck):
    numbers = {
        token
        for text in prose_and_notes(deck)
        for token in NUMBER.findall(CLOCK.sub("", text))
    }
    assert numbers - MEASURED.keys() - SET.keys() == set()


def test_every_measured_number_names_a_claim_test_that_exists():
    claims = (Path(__file__).parent / "test_measured_claims.py").read_text()
    for test_name in MEASURED.values():
        assert f"def {test_name}(" in claims, test_name


@pytest.mark.slow
def test_the_owners_placeholders_are_visible(deck):
    placeholders = [
        shape.text_frame.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.name == "placeholder"
    ]
    for owed in (
        "[photo]",
        "[bio]",
        "[email]",
        "[survey link]",
        "[Zenodo DOI",
        "[owner: the AI-use panel",
        "[owner: one line on The Individuating Membrane",
    ):
        assert any(text.startswith(owed) for text in placeholders), owed


@pytest.mark.slow
def test_claude_is_named_only_in_the_owners_disclosure(deck):
    naming = [text for text in prose_and_notes(deck) if "Claude" in text]
    assert naming[0] == build_deck.DISCLOSURE
    assert len(naming) == 2  # the caption and the Q&A answer in its notes
    assert naming[1].endswith("the answer is: some of it.")


def snapshot(presentation):
    slides = []
    for slide in presentation.slides:
        shapes = [
            (shape.name, shape.left, shape.top, shape.width, shape.height,
             shape.text_frame.text if shape.has_text_frame else
             hashlib.sha256(shape.image.blob).hexdigest()
             if shape.name.startswith("picture") else None)
            for shape in slide.shapes
        ]  # fmt: skip
        slides.append((shapes, slide.notes_slide.notes_text_frame.text))
    return slides


def test_links_are_drawn_in_ink_not_the_templates_dark_blue(deck):
    theme = deck.slide_master.part.part_related_by(RELATIONSHIP_TYPE.THEME)
    scheme = ET.fromstring(theme.blob).find(f".//{{{DRAWING}}}clrScheme")
    assert scheme is not None
    for name in ("hlink", "folHlink"):
        colour = scheme.find(f"{{{DRAWING}}}{name}/{{{DRAWING}}}srgbClr")
        assert colour is not None
        assert colour.get("val") == plot.LIGHT.ink.lstrip("#")


@pytest.mark.slow
def test_the_committed_deck_is_the_one_the_script_builds(deck):
    committed = Presentation(str(build_deck.DECK))
    assert snapshot(committed) == snapshot(deck)
